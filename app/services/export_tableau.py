"""Export d'un tableau tel qu'affiché à l'écran : Excel ou PDF.

L'écran envoie ses colonnes et ses lignes déjà filtrées, triées et mises en
forme ; ce module ne fait que les poser dans un fichier. Ce qu'on télécharge
est donc exactement ce qu'on voit, sans second calcul côté serveur.

Une ligne est `{"cellules": [...], "groupe": bool}`. Une ligne de groupe
(titre de catégorie, matière dans les vues fournisseur) occupe toute la
largeur et ne porte que sa première cellule.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Dict, List


def _texte(v: Any) -> str:
    return "" if v is None else str(v)


def classeur(titre: str, sous_titre: str, colonnes: List[str],
             lignes: List[Dict[str, Any]], genere_le: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (titre or "Export")[:31]
    nb = max(1, len(colonnes))
    ws.append([titre])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([(sous_titre + " · " if sous_titre else "") + "généré le " + genere_le.strftime("%d/%m/%Y %H:%M")])
    ws["A2"].font = Font(italic=True, color="64748B")
    ws.append([])
    ws.append(colonnes)
    entete = 4
    for c in range(1, nb + 1):
        cell = ws.cell(row=entete, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0891B2")
    fond_groupe = PatternFill("solid", fgColor="E0F7FA")
    for ligne in lignes:
        cellules = list(ligne.get("cellules") or [])
        ws.append(cellules[:1] if ligne.get("groupe") else cellules)
        if ligne.get("groupe"):
            r = ws.max_row
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=nb)
            ws.cell(row=r, column=1).font = Font(bold=True)
            ws.cell(row=r, column=1).fill = fond_groupe
    for c in range(1, nb + 1):
        larg = max([len(_texte(colonnes[c - 1])) if c - 1 < len(colonnes) else 0]
                   + [len(_texte((l.get("cellules") or [None] * nb)[c - 1]))
                      for l in lignes if not l.get("groupe") and len(l.get("cellules") or []) >= c])
        ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, larg + 2))
        for r in range(entete + 1, ws.max_row + 1):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = ws.cell(row=entete + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf(titre: str, sous_titre: str, colonnes: List[str], lignes: List[Dict[str, Any]],
        genere_le: datetime, genere_par: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    accent, accent_bg = colors.HexColor("#0891b2"), colors.HexColor("#e0f7fa")
    muted, text, border = colors.HexColor("#64748b"), colors.HexColor("#0f172a"), colors.HexColor("#e2e8f0")
    st_titre = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, textColor=text, leading=20)
    st_sous = ParagraphStyle("s", fontName="Helvetica", fontSize=9, textColor=muted, leading=12)
    st_c = ParagraphStyle("c", fontName="Helvetica", fontSize=8, textColor=text, leading=10)
    st_h = ParagraphStyle("h", parent=st_c, fontName="Helvetica-Bold", textColor=colors.white)
    st_g = ParagraphStyle("g", parent=st_c, fontName="Helvetica-Bold")

    page = landscape(A4)
    largeur = page[0] - 20 * mm
    nb = max(1, len(colonnes))

    def p(v, st=st_c):
        return Paragraph(_esc(_texte(v)), st)

    data = [[p(c, st_h) for c in colonnes]]
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("GRID", (0, 0), (-1, -1), 0.4, border),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for ligne in lignes:
        cellules = list(ligne.get("cellules") or [])
        if ligne.get("groupe"):
            r = len(data)
            data.append([p(cellules[0] if cellules else "", st_g)] + [""] * (nb - 1))
            style += [("SPAN", (0, r), (-1, r)), ("BACKGROUND", (0, r), (-1, r), accent_bg)]
        else:
            cellules += [""] * (nb - len(cellules))
            data.append([p(v) for v in cellules[:nb]])

    def pied(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(10 * mm, 8 * mm, titre + " · généré le "
                          + genere_le.strftime("%d/%m/%Y à %H:%M") + (f" par {genere_par}" if genere_par else ""))
        canvas.drawRightString(page[0] - 10 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=14 * mm, title=titre, author="MySifa")
    tableau = Table(data, colWidths=[largeur / nb] * nb, repeatRows=1)
    tableau.setStyle(TableStyle(style))
    story = [Paragraph(_esc(titre), st_titre)]
    if sous_titre:
        story.append(Paragraph(_esc(sous_titre), st_sous))
    story += [Spacer(1, 6), tableau]
    doc.build(story, onFirstPage=pied, onLaterPages=pied)
    return buf.getvalue()


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
