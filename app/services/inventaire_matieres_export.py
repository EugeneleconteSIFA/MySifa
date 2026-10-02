"""
Export de l'inventaire matières (MyStock › Inventaire matière) : Excel et PDF.

Les deux formats partent des MÊMES lignes (`stock.matieres_inventaire_donnees`),
pour qu'un chiffre lu dans le PDF soit celui du classeur.

Ce que l'export doit permettre, et qui manquait à l'ancien fichier (une feuille,
une ligne par référence, laizes et emplacements entassés dans deux cellules) :

- compter : une ligne par laize pour les bobines, avec la case « compté » et
  l'écart calculé ; une ligne par emplacement, triée par emplacement, pour
  faire le tour du magasin dans l'ordre ;
- lire le stock comme on le consomme : les bobines en mètres (bobines ×
  métrage standard de la fiche) à côté du nombre de bobines ;
- ne pas compter deux fois : les non-conformités (emplacements « NC … ») sont
  hors stock disponible, signalées comme telles ;
- ne garder que l'utile : une laize sans stock ne fait pas une ligne.

Le PDF est la feuille de comptage du magasin : par catégorie, une ligne par
référence (par laize pour les bobines), une case vide à remplir à la main.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any, Dict, List, Optional

STATUTS = {"rouge": "À faire", "orange": "Bientôt", "vert": "À jour"}
_RE_NC = re.compile(r"^NC\b")


def est_nc(emplacement: Any) -> bool:
    """Un emplacement « NC … » porte une non-conformité, hors stock disponible."""
    return bool(_RE_NC.match(str(emplacement or "").strip()))


def _fmt_n(v: Optional[float], dec: int = 3) -> str:
    if v is None:
        return "—"
    f = round(float(v), dec)
    if abs(f - round(f)) < 10 ** -dec:
        txt = f"{int(round(f)):,}".replace(",", " ")
    else:
        txt = f"{f:,.{dec}f}".rstrip("0").rstrip(".").replace(",", " ").replace(".", ",")
    return txt


def _abr(unite: str) -> str:
    return {"bobine": "bob.", "palette": "pal.", "kg": "kg", "unite": "u.", "pile": "piles"}.get(unite or "", unite or "")


def _date_fr(d: Optional[datetime]) -> str:
    return d.strftime("%d/%m/%Y") if d else "Jamais"


# ── Excel ────────────────────────────────────────────────────────────────────

def classeur(matieres: List[Dict[str, Any]], titre_filtre: str, genere_le: datetime) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1E293B")
    head_font = Font(bold=True, color="FFFFFF")
    saisie_fill = PatternFill("solid", fgColor="FEF9C3")
    nc_font = Font(color="B91C1C")

    wb = Workbook()

    def feuille(ws, titre, entetes, largeurs, lignes, col_theo=None, formats=None, gras=None):
        ws.title = titre[:31]
        ws.append([f"Inventaire matières — {titre}"])
        ws.append([f"Généré le {genere_le.strftime('%d/%m/%Y à %H:%M')}" + (f" · {titre_filtre}" if titre_filtre else "")])
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"].font = Font(color="64748B", size=9)
        cols = list(entetes) + (["Quantité comptée", "Écart"] if col_theo else [])
        ws.append(cols)
        h = ws.max_row
        for i in range(1, len(cols) + 1):
            c = ws.cell(row=h, column=i)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        n = len(entetes)
        for ligne in lignes:
            ws.append(list(ligne) + ([None, None] if col_theo else []))
            r = ws.max_row
            if col_theo:
                lt, lc = get_column_letter(col_theo), get_column_letter(n + 1)
                ws.cell(row=r, column=n + 1).fill = saisie_fill
                ws.cell(row=r, column=n + 2).value = f'=IF({lc}{r}="","",{lc}{r}-{lt}{r})'
            for col, fmt in (formats or {}).items():
                ws.cell(row=r, column=col).number_format = fmt
            if gras and gras(ligne):
                for i in range(1, n + 1):
                    ws.cell(row=r, column=i).font = nc_font
        for i, w in enumerate(list(largeurs) + ([16, 12] if col_theo else []), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=h + 1, column=1)
        ws.auto_filter.ref = f"A{h}:{get_column_letter(len(cols))}{max(ws.max_row, h)}"

    NB = "#,##0.###"
    M = "#,##0"

    # 1. Synthèse par catégorie
    synth: Dict[tuple, Dict[str, Any]] = {}
    for m in matieres:
        k = (m["categorie_label"], m["sous_section"] or "", m["unite"])
        s = synth.setdefault(k, {"n": 0, "stock": 0.0, "metres": None, "nc": 0.0,
                                 "rouge": 0, "orange": 0, "vert": 0})
        s["n"] += 1
        s["stock"] += m["stock"]
        if m["metres"] is not None:
            s["metres"] = (s["metres"] or 0) + m["metres"]
        s["nc"] += m["nc"]
        s[m["statut"]] = s.get(m["statut"], 0) + 1
    feuille(wb.active, "Synthèse",
            ["Catégorie", "Sous-section", "Références", "Stock", "Unité", "Métrage (m)",
             "Hors stock (NC)", "À faire", "Bientôt", "À jour"],
            [16, 20, 11, 13, 9, 15, 15, 9, 9, 9],
            [[k[0], k[1], v["n"], round(v["stock"], 3), _abr(k[2]),
              round(v["metres"]) if v["metres"] is not None else None, round(v["nc"], 3) or None,
              v["rouge"], v["orange"], v["vert"]] for k, v in sorted(synth.items())],
            formats={4: NB, 6: M, 7: NB})

    # 2. Par référence
    feuille(wb.create_sheet(), "Par référence",
            ["Catégorie", "Sous-section", "Libellé commercial", "Référence", "Stock", "Unité",
             "Métrage (m)", "Laizes en stock", "Hors stock (NC)", "Dernier inventaire",
             "Inventorié par", "Statut"],
            [13, 16, 38, 22, 11, 8, 13, 26, 13, 15, 18, 10],
            [[m["categorie_label"], m["sous_section"] or "", m["designation"], m["reference_si_diff"],
              round(m["stock"], 3), _abr(m["unite"]),
              round(m["metres"]) if m["metres"] is not None else None,
              ", ".join(_fmt_n(l["laize_mm"], 1) for l in m["laizes"] if l["stock"]) or "",
              round(m["nc"], 3) or None, m["derniere"] or "Jamais", m["operateur"] or "",
              STATUTS.get(m["statut"], "")] for m in matieres],
            col_theo=5, formats={5: NB, 7: M, 9: NB, 10: "dd/mm/yyyy"})

    # 3. Par laize (bobines) — la feuille de comptage des frontaux et glassines
    lignes_lz = []
    for m in matieres:
        for l in m["laizes"]:
            if not l["stock"] and not l["emplacements"]:
                continue
            lignes_lz.append([m["categorie_label"], m["sous_section"] or "", m["designation"],
                              l["laize_mm"], round(l["stock"], 3),
                              round(l["metres"]) if l["metres"] is not None else None,
                              " · ".join(f"{e['emplacement']} : {_fmt_n(e['quantite'])}"
                                         for e in l["emplacements"])])
    feuille(wb.create_sheet(), "Par laize",
            ["Catégorie", "Sous-section", "Libellé commercial", "Laize (mm)", "Stock (bob.)",
             "Métrage (m)", "Emplacements"],
            [13, 16, 38, 11, 12, 13, 40], lignes_lz,
            col_theo=5, formats={4: "0.#", 5: NB, 6: M})

    # 4. Par emplacement — le tour du magasin, dans l'ordre des emplacements
    lignes_e = []
    for m in matieres:
        for e in m["emplacements"]:
            lignes_e.append([e["emplacement"], m["categorie_label"], m["designation"],
                             e["laize_mm"], round(e["quantite"], 3), _abr(m["unite"]),
                             "Non-conformité" if e["nc"] else "", e["maj_le"], e["maj_par"] or ""])
    lignes_e.sort(key=lambda r: (str(r[0]), str(r[2]), r[3] or 0))
    feuille(wb.create_sheet(), "Par emplacement",
            ["Emplacement", "Catégorie", "Libellé commercial", "Laize (mm)", "Quantité", "Unité",
             "Hors stock", "Mis à jour le", "Par"],
            [16, 13, 38, 11, 11, 8, 15, 15, 18], lignes_e,
            col_theo=5, formats={4: "0.#", 5: NB, 8: "dd/mm/yyyy"},
            gras=lambda r: bool(r[6]))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

def pdf(matieres: List[Dict[str, Any]], titre_filtre: str, genere_le: datetime,
        genere_par: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    accent, accent_bg = colors.HexColor("#0891b2"), colors.HexColor("#e0f7fa")
    muted, text, border = colors.HexColor("#64748b"), colors.HexColor("#0f172a"), colors.HexColor("#e2e8f0")
    rouge, orange, vert = colors.HexColor("#b91c1c"), colors.HexColor("#c2410c"), colors.HexColor("#15803d")
    st_titre = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=text, leading=21)
    st_sous = ParagraphStyle("s", fontName="Helvetica", fontSize=9, textColor=muted, leading=12)
    st_h2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, textColor=text,
                           spaceBefore=10, spaceAfter=4)
    st_c = ParagraphStyle("c", fontName="Helvetica", fontSize=8, textColor=text, leading=10)
    st_cb = ParagraphStyle("cb", parent=st_c, fontName="Helvetica-Bold")
    st_cm = ParagraphStyle("cm", parent=st_c, textColor=muted, fontSize=7, leading=9)
    st_r = ParagraphStyle("r", parent=st_c, alignment=2)

    largeur = landscape(A4)[0] - 20 * mm

    def pied(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(muted)
        canvas.drawString(10 * mm, 8 * mm, "Inventaire matières · généré le "
                          + genere_le.strftime("%d/%m/%Y à %H:%M") + (f" par {genere_par}" if genere_par else ""))
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=14 * mm,
                            title="Inventaire matières", author="MySifa")
    story: list = [
        Paragraph("Inventaire matières", st_titre),
        Paragraph("Feuille de comptage — " + genere_le.strftime("%d/%m/%Y")
                  + (" · " + titre_filtre if titre_filtre else ""), st_sous),
        Paragraph("Bobines : une ligne par laize, métrage = bobines × métrage standard de la fiche. "
                  "Les non-conformités (emplacements NC) sont hors stock disponible.", st_sous),
        Spacer(1, 6),
    ]

    # Synthèse
    synth: Dict[str, Dict[str, Any]] = {}
    for m in matieres:
        s = synth.setdefault(m["categorie_label"], {"n": 0, "stock": 0.0, "metres": None,
                                                    "unite": m["unite"], "rouge": 0, "orange": 0, "vert": 0,
                                                    "mixte": False})
        s["n"] += 1
        if s["unite"] != m["unite"]:
            s["mixte"] = True
        s["stock"] += m["stock"]
        if m["metres"] is not None:
            s["metres"] = (s["metres"] or 0) + m["metres"]
        s[m["statut"]] += 1
    rows = [["Catégorie", "Références", "Stock", "Métrage", "À faire", "Bientôt", "À jour"]]
    for cat, s in sorted(synth.items()):
        rows.append([cat, str(s["n"]),
                     "—" if s["mixte"] else f"{_fmt_n(s['stock'])} {_abr(s['unite'])}",
                     f"{_fmt_n(s['metres'], 0)} m" if s["metres"] is not None else "—",
                     str(s["rouge"]), str(s["orange"]), str(s["vert"])])
    t = Table(rows, colWidths=[60 * mm, 25 * mm, 40 * mm, 40 * mm, 22 * mm, 22 * mm, 22 * mm], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), accent_bg), ("TEXTCOLOR", (0, 0), (-1, 0), accent),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, 0), 1, accent),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, border),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [t, Spacer(1, 4)]

    # Feuilles de comptage, par catégorie
    cw = [62 * mm, 18 * mm, 24 * mm, 26 * mm, 62 * mm, 22 * mm, 14 * mm, 49 * mm]
    cw[-1] = largeur - sum(cw[:-1])
    entete = ["Libellé commercial", "Laize", "Stock", "Métrage", "Emplacements",
              "Dernier inv.", "Statut", "Compté"]
    par_cat: Dict[str, List[Dict[str, Any]]] = {}
    for m in matieres:
        par_cat.setdefault(m["categorie_label"] + (" · " + m["sous_section"] if m["sous_section"] else ""),
                           []).append(m)
    for cat in sorted(par_cat):
        data = [[Paragraph(h, ParagraphStyle("h", parent=st_cb, textColor=accent)) for h in entete]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), accent_bg), ("LINEBELOW", (0, 0), (-1, 0), 1, accent),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]
        for m in par_cat[cat]:
            statut_c = {"rouge": rouge, "orange": orange, "vert": vert}.get(m["statut"], muted)
            titre = [Paragraph(_esc(m["designation"]), st_cb)]
            if m["reference_si_diff"]:
                titre.append(Paragraph("Réf. " + _esc(m["reference_si_diff"]), st_cm))
            if m["nc"]:
                titre.append(Paragraph(f"Hors stock : {_fmt_n(m['nc'])} {_abr(m['unite'])} en non-conformité",
                                       ParagraphStyle("nc", parent=st_cm, textColor=rouge)))
            lz = [l for l in m["laizes"] if l["stock"] or l["emplacements"]] if m["laizee"] else []
            sous = lz or [None]
            debut = len(data)
            for i, l in enumerate(sous):
                if l is None:
                    stock, metres, empl, laize = m["stock"], m["metres"], m["emplacements"], "—"
                else:
                    stock, metres, empl, laize = l["stock"], l["metres"], l["emplacements"], _fmt_n(l["laize_mm"], 1)
                data.append([
                    titre if i == 0 else "",
                    Paragraph(laize, st_r),
                    Paragraph(f"{_fmt_n(stock)} {_abr(m['unite'])}", st_r),
                    Paragraph(f"{_fmt_n(metres, 0)} m" if metres is not None else "—", st_r),
                    Paragraph(" · ".join(_esc(f"{e['emplacement']} : {_fmt_n(e['quantite'])}")
                                         for e in empl) or "—", st_c),
                    Paragraph(_date_fr(m["derniere"]), st_c) if i == 0 else "",
                    Paragraph(STATUTS.get(m["statut"], ""), ParagraphStyle("st", parent=st_cb, textColor=statut_c))
                    if i == 0 else "",
                    "",
                ])
            fin = len(data) - 1
            if fin > debut:
                for col in (0, 5, 6):
                    style.append(("SPAN", (col, debut), (col, fin)))
                style.append(("VALIGN", (0, debut), (0, fin), "TOP"))
            style.append(("LINEBELOW", (0, fin), (-1, fin), 0.5, border))
            for r in range(debut, fin + 1):
                style.append(("BOX", (7, r), (7, r), 0.6, muted))
        tbl = Table(data, colWidths=cw, repeatRows=1)
        tbl.setStyle(TableStyle(style))
        story.append(KeepTogether([Paragraph(_esc(cat) + f" — {len(par_cat[cat])} référence"
                                             + ("s" if len(par_cat[cat]) > 1 else ""), st_h2)]))
        story.append(tbl)
    if not matieres:
        story.append(Paragraph("Aucune matière pour ce filtre.", st_sous))
    doc.build(story, onFirstPage=pied, onLaterPages=pied)
    return buf.getvalue()


def _esc(s: Any) -> str:
    return (str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
