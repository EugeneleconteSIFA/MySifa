# -*- coding: utf-8 -*-
"""Registre FSC des approvisionnements : import, éligibilité, journal, export.

Ce que ce test protège :

1. **L'éligibilité ne tranche pas trop tôt.** Une date d'expiration inconnue
   donne « à vérifier », jamais « non » ; un BL et une facture qui divergent
   donnent leur propre verdict au lieu de se ranger dans « non ».
2. **La quantité est en mètres linéaires.** `cua` annonce des m² sur les
   bobines ; croire `cua` fausserait tous les volumes du bilan d'un facteur égal
   à la laize.
3. **L'import est idempotent** et ne réécrit jamais une ligne déjà saisie : une
   resynchro quotidienne ne doit pas effacer le travail de la veille.
4. **Le rapprochement fournisseur se tait quand il hésite.** « FRIMPEKS LTD »
   peut être Italy, UK ou Turkey : la ligne entre au registre non rattachée
   plutôt que rattachée au hasard.
5. **Toute correction laisse une trace.**

Lancer : python3 tests/test_fsc_registre.py
"""

from __future__ import annotations

import importlib.util
import io
import os
import sqlite3
import sys

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)

from app.services import fsc_registre as R  # noqa: E402

ECHECS: list[str] = []


def verifie(condition, libelle: str) -> None:
    if condition:
        print("  ok   %s" % libelle)
    else:
        print("  ECHEC %s" % libelle)
        ECHECS.append(libelle)


def egal(obtenu, attendu, libelle: str) -> None:
    verifie(obtenu == attendu, "%s (obtenu %r, attendu %r)" % (libelle, obtenu, attendu))


# ── Bases minimales ────────────────────────────────────────────────────────
SCHEMA_MYSIFA = """
CREATE TABLE fournisseurs_fsc (
    id INTEGER PRIMARY KEY, nom TEXT NOT NULL, licence TEXT, certificat TEXT,
    has_fsc INTEGER NOT NULL DEFAULT 1, actif INTEGER NOT NULL DEFAULT 1,
    fsc_date_expiration TEXT, rvgi_numero INTEGER, rvgi_rs TEXT, rvgi_lie_le TEXT,
    updated_at TEXT);
"""
SCHEMA_ERP = """
CREATE TABLE lif_ligne (id INTEGER PRIMARY KEY, corbeille INTEGER DEFAULT 0, numero INTEGER,
    ligne INTEGER, qte REAL, amjl TEXT, dtem TEXT, ref TEXT, fac_no INTEGER, fac_lg INTEGER);
CREATE TABLE cdf_ligne (id INTEGER PRIMARY KEY, corbeille INTEGER DEFAULT 0, numero INTEGER,
    ligne INTEGER, type INTEGER, code1 TEXT, code2 TEXT, code3 TEXT, des1 TEXT, cua TEXT);
CREATE TABLE cdf_entete (id INTEGER PRIMARY KEY, corbeille INTEGER DEFAULT 0, numero INTEGER,
    numfou INTEGER, rs TEXT);
CREATE TABLE mat_mat (id INTEGER PRIMARY KEY, corbeille INTEGER DEFAULT 0, code1 TEXT, code2 TEXT,
    type INTEGER, libc1 TEXT, ref TEXT, libt2 TEXT);
"""


def base_mysifa():
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    c.executescript(SCHEMA_MYSIFA)
    c.executemany(
        "INSERT INTO fournisseurs_fsc (id,nom,licence,certificat,fsc_date_expiration,rvgi_numero) "
        "VALUES (?,?,?,?,?,?)",
        [
            (1, "Likexin", "FSC-C128270", "ESTS-COC-242264", "2030-02-10", None),
            (2, "Burgo / Mosaico", "FSC-C004657", "SGSCH-COC-002122", "2028-02-11", None),
            (3, "Frimpeks Italy", "FSC-C164660", "INT-COC-001611", "2031-02-24", None),
            (4, "Frimpeks UK", "FSC-C160714", "INT-COC-002144", "2030-10-11", None),
            (5, "Lefrancq", "FSC-C135176", "FCBA-COC-000478", None, None),
            (6, "Kanzan", "FSC-C007179", "TUVDC-COC-100605", "2024-01-01", 4242),
        ],
    )
    spec = importlib.util.spec_from_file_location(
        "mig_fsc_registre",
        os.path.join(RACINE, "app", "core", "migrations", "2026_09_14_fsc_registre_appro.py"))
    mig = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mig)
    mig.appliquer(c)
    mig.appliquer(c)  # rejouable
    return c


def base_erp():
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    c.executescript(SCHEMA_ERP)
    c.executemany(
        "INSERT INTO cdf_entete (numero,numfou,rs) VALUES (?,?,?)",
        [(9001, 101, "SHENZHEN LIKEXIN INDUSTRIAL Co."), (9002, 102, "MOSAICO"),
         (9003, 103, "FRIMPEKS LTD"), (9004, 4242, "KANZAN SPEZIALPAPIERE")],
    )
    c.executemany(
        "INSERT INTO cdf_ligne (numero,ligne,type,code1,code2,code3,des1,cua) VALUES (?,?,?,?,?,?,?,?)",
        [(9001, 1, 7, "1183", "0004", "470", "70gsm Direct Thermal", "10"),
         (9001, 2, 4, "1183", "0001", "440", "60g Yellow glassine", "10"),
         (9002, 1, 3, "629", "0002", "333", "Velin permanent", "M²"),
         (9003, 1, 3, "700", "0001", "500", "Complexe", "10"),
         (9004, 1, 7, "574", "0010", "510", "Thermique Pro", "10"),
         # Hors périmètre : adhésif (type 9), il ne doit pas entrer au registre.
         (9004, 2, 9, "800", "0001", "", "Adhésif", "KG")],
    )
    c.executemany(
        "INSERT INTO lif_ligne (id,numero,ligne,qte,amjl,dtem,ref,fac_no,fac_lg) VALUES (?,?,?,?,?,?,?,?,?)",
        [(10, 9001, 1, 10000, "2025-12-01 08:00", "2025-12-01 08:00", "BL-AVANT", 1, 1),
         (11, 9001, 1, 84210, "2026-07-31 10:00", "2026-08-02 09:00", "FR20260609", 3610, 1),
         (12, 9001, 2, 218340, "2026-07-31 10:00", "2026-08-02 09:00", "FR20260609", 3610, 2),
         (13, 9002, 1, 11600, "2026-01-13 07:00", "2026-01-13 07:00", "M26001289", 3617, 1),
         (14, 9003, 1, 5000, "2026-02-02 07:00", "2026-02-02 07:00", "FRIM-1", 3620, 1),
         (15, 9004, 1, 32800, "2026-01-08 07:00", "2026-01-08 07:00", "252631", 3611, 1),
         (16, 9004, 2, 900, "2026-01-08 07:00", "2026-01-08 07:00", "252631", 3611, 2)],
    )
    c.executemany(
        "INSERT INTO mat_mat (code1,code2,type,libc1,ref,libt2) VALUES (?,?,?,?,?,?)",
        [("1183", "0004", 5, "Thermique Pro TOP 70g", "LKX-70", "Roll 12.000 ml"),
         ("1183", "0004", 2, "NE PAS PRENDRE — fiche glassine", "X", ""),
         ("1183", "0001", 2, "Glassine 60g", "LKX-GL", "Roll 18.000 ml")],
    )
    return c


# ── 1. Éligibilité ─────────────────────────────────────────────────────────
print("\n--- 1. règle d'éligibilité ---")
BASE = {"certificat_statut": "valide", "allegation_bl": "fsc_mix_credit",
        "allegation_facture": "fsc_mix_credit", "code_certificat_present": 1,
        "pourcentage": None}
egal(R.evaluer_eligibilite(BASE), "oui", "les quatre conditions réunies")
egal(R.evaluer_eligibilite({**BASE, "certificat_statut": "inconnu"}), "a_verifier",
     "date d'expiration inconnue : à vérifier, pas « non »")
egal(R.evaluer_eligibilite({**BASE, "certificat_statut": None}), "a_verifier",
     "fournisseur non rattaché : à vérifier")
egal(R.evaluer_eligibilite({**BASE, "certificat_statut": "expire"}), "non",
     "certificat expiré à la date du BL")
egal(R.evaluer_eligibilite({**BASE, "certificat_statut": "non_certifie"}), "non",
     "fournisseur non certifié")
egal(R.evaluer_eligibilite({**BASE, "allegation_facture": "fsc_mix_pct"}), "ecart_bl_facture",
     "BL et facture divergents")
egal(R.evaluer_eligibilite({**BASE, "allegation_bl": "aucune", "allegation_facture": "aucune"}),
     "non", "aucune allégation facturée")
egal(R.evaluer_eligibilite({**BASE, "allegation_bl": "fsc_controlled_wood",
                            "allegation_facture": "fsc_controlled_wood"}), "non",
     "Controlled Wood ne rend pas un produit vendable avec allégation")
egal(R.evaluer_eligibilite({**BASE, "code_certificat_present": 0}), "non",
     "code de certificat absent des documents")
egal(R.evaluer_eligibilite({**BASE, "code_certificat_present": None}), "a_verifier",
     "code de certificat pas encore contrôlé")
egal(R.evaluer_eligibilite({**BASE, "allegation_bl": "fsc_mix_pct",
                            "allegation_facture": "fsc_mix_pct"}), "a_verifier",
     "allégation en % sans pourcentage")
egal(R.evaluer_eligibilite({**BASE, "allegation_bl": "fsc_mix_pct",
                            "allegation_facture": "fsc_mix_pct", "pourcentage": 70}), "oui",
     "allégation en % avec pourcentage")
verifie(R.label_autorise("fsc_mix_pct", 70) and not R.label_autorise("fsc_mix_pct", 50),
        "label FSC : seuil de 70 % sur les allégations en pourcentage")
verifie(R.label_autorise("fsc_100") and not R.label_autorise("fsc_controlled_wood"),
        "label FSC : autorisé sur FSC 100 %, jamais sur Controlled Wood")


# ── 2. Rapprochement fournisseur ───────────────────────────────────────────
print("\n--- 2. rapprochement fournisseur ---")
conn = base_mysifa()
par_numero, par_nom = R.index_fournisseurs(conn)
egal((R.trouver_fournisseur(101, "SHENZHEN LIKEXIN INDUSTRIAL Co.", par_numero, par_nom) or {}).get("nom"),
     "Likexin", "nom d'usage contenu dans la raison sociale")
egal((R.trouver_fournisseur(102, "MOSAICO", par_numero, par_nom) or {}).get("nom"),
     "Burgo / Mosaico", "raison sociale contenue dans le nom d'usage")
egal(R.trouver_fournisseur(103, "FRIMPEKS LTD", par_numero, par_nom), None,
     "deux branches possibles : aucune réponse")
egal((R.trouver_fournisseur(4242, "PEU IMPORTE LE NOM", par_numero, par_nom) or {}).get("nom"),
     "Kanzan", "le numéro RVGI l'emporte sur le nom")


# ── 3. Import ──────────────────────────────────────────────────────────────
print("\n--- 3. import depuis le miroir ---")
erp = base_erp()
egal(R.date_entree(conn), None, "aucune date d'entrée au départ")
egal(R.importer(conn, erp)["ajoutees"], 0, "sans date d'entrée, rien ne s'importe")
R.definir_date_entree(conn, "2026-01-01", "test")
bilan = R.importer(conn, erp, "test")
egal(bilan["ajoutees"], 5, "5 lignes importées (l'adhésif et la réception de 2025 restent dehors)")
egal(bilan["sans_fournisseur"], 1, "la ligne FRIMPEKS LTD entre non rattachée")
egal(R.importer(conn, erp, "test")["ajoutees"], 0, "deuxième passage : rien de neuf")

lignes = {l["lif_id"]: l for l in R.lister(conn, limite=100)}
egal(len(lignes), 5, "cinq lignes au registre")
verifie(10 not in lignes, "la réception antérieure à la date d'entrée est écartée")
verifie(16 not in lignes, "l'adhésif (type 9) n'entre pas au registre")
l11 = lignes[11]
egal(l11["quantite_ml"], 84210, "la quantité reste en mètres linéaires")
egal(l11["quantite_m2"], round(84210 * 470 / 1000.0, 2), "m² = ml × laize / 1000")
egal(l11["libelle_matiere"], "Thermique Pro TOP 70g",
     "fiche matière jointe sur le triplet (type = type d'achat - 2)")
egal(l11["type_matiere"], "frontal", "type d'article traduit en catégorie MySifa")
egal(l11["certificat_statut"], "valide", "certificat jugé à la date du BL")
egal(lignes[15]["certificat_statut"], "expire", "certificat expiré à la date du BL")
egal(lignes[15]["eligible"], "non", "et la ligne n'est pas éligible")
egal(lignes[14]["certificat_statut"], None, "fournisseur non rattaché : pas de verdict")
egal(lignes[14]["eligible"], "a_verifier", "et la ligne reste à vérifier")


# ── 4. Saisie, propagation au BL, journal ─────────────────────────────────
print("\n--- 4. saisie et journal ---")
saisie = {"allegation_bl": "fsc_mix_credit", "allegation_facture": "fsc_mix_credit",
          "code_certificat_present": 1, "num_facture_fournisseur": "F-2026-118"}
maj = R.mettre_a_jour(conn, lignes[11]["id"], saisie, "Fatiha")
egal(maj["eligible"], "oui", "la saisie rend la ligne éligible")
egal(maj["controle_par"], "Fatiha", "le contrôleur est enregistré")
egal(R.appliquer_au_bl(conn, lignes[11]["id"], "Fatiha"), 1,
     "la saisie se propage à l'autre ligne du même BL")
apres = {l["lif_id"]: l for l in R.lister(conn, limite=100)}
egal(apres[12]["num_facture_fournisseur"], "F-2026-118", "et le numéro de facture avec")
egal(apres[12]["eligible"], "oui", "la ligne propagée est éligible elle aussi")
verifie(apres[12]["observations"] == "", "l'observation, elle, ne se propage pas")

j = R.journal(conn, lignes[11]["id"])
verifie(any(e["champ"] == "allegation_facture" and e["nouvelle_valeur"] == "fsc_mix_credit" for e in j),
        "le journal garde la valeur saisie")
verifie(any(e["champ"] == "eligible" and e["ancienne_valeur"] == "a_verifier" for e in j),
        "et le changement de verdict, avec son avant")
n_avant = len(R.journal(conn, lignes[11]["id"]))
R.mettre_a_jour(conn, lignes[11]["id"], {"allegation_bl": "fsc_mix_credit"}, "Fatiha")
egal(len(R.journal(conn, lignes[11]["id"])), n_avant, "réécrire la même valeur n'écrit pas au journal")

for champs, motif in (({"allegation_bl": "fsc_premium"}, "allégation hors liste"),
                      ({"pourcentage": 150}, "pourcentage hors plage"),
                      ({"etiquette_posee": "bleue"}, "étiquette hors liste")):
    try:
        R.mettre_a_jour(conn, lignes[11]["id"], champs, "Fatiha")
        verifie(False, "refus attendu : %s" % motif)
    except ValueError:
        verifie(True, "refus attendu : %s" % motif)

R.mettre_a_jour(conn, lignes[13]["id"], {"observations": "Reste à demander la facture"}, "Fatiha")
egal({l["lif_id"]: l for l in R.lister(conn, limite=100)}[13]["eligible"], "a_verifier",
     "une observation seule ne rend pas une ligne éligible")


# ── 5. Volumes et export ───────────────────────────────────────────────────
print("\n--- 5. volumes et export ---")
toutes = R.lister(conn, limite=100)
vol = R.volumes_par_allegation(toutes)
egal([v["code"] for v in vol], ["fsc_mix_credit"], "une seule allégation éligible")
egal(vol[0]["lignes"], 2, "les deux lignes du BL comptent")
egal(vol[0]["m2"], round((84210 * 470 + 218340 * 440) / 1000.0, 2), "m² cumulés")
egal(R.stats(toutes)["sans_fournisseur"], 1, "le non rattaché reste compté à part")
egal(len(R.lister(conn, debut="2026-01-01", fin="2026-01-31")), 2, "filtre de période")
egal(len(R.lister(conn, eligible="oui")), 2, "filtre d'éligibilité")
egal(len(R.lister(conn, certifies_seuls=True)), 4, "filtre fournisseurs certifiés")
egal(len(R.lister(conn, recherche="FR20260609")), 2, "recherche par n° de BL")

contenu = R.export_xlsx(toutes, "2026-01-01", "2026-12-31")
verifie(contenu[:2] == b"PK", "le classeur est un vrai fichier xlsx")
try:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenu))
    egal(wb.sheetnames, ["Approvisionnements", "Volumes par allégation"], "deux feuilles")
    ws = wb["Approvisionnements"]
    egal(ws.cell(row=4, column=1).value, "Date de réception", "en-têtes en ligne 4")
    egal(ws.max_row, 4 + len(toutes), "une ligne par réception")
except ImportError:
    print("  (openpyxl absent : relecture du classeur non vérifiée)")


# ── 6. Rattachement manuel ─────────────────────────────────────────────────
print("\n--- 6. rattachement d'un tiers RVGI ---")
egal(R.rattacher_fournisseur(conn, 103, 3, "FRIMPEKS LTD", "Eugène"), 1,
     "la ligne non rattachée est reprise")
reprise = {l["lif_id"]: l for l in R.lister(conn, limite=100)}[14]
egal(reprise["fournisseur_id"], 3, "elle porte maintenant la fiche choisie")
egal(reprise["certificat_statut"], "valide", "et le verdict du certificat est calculé")
egal(conn.execute("SELECT rvgi_numero FROM fournisseurs_fsc WHERE id=3").fetchone()[0], 103,
     "le numéro RVGI est mémorisé sur la fiche")
verifie(any(e["champ"] == "fournisseur_id" for e in R.journal(conn, reprise["id"])),
        "le rattachement est journalisé")


print()
if ECHECS:
    print("%d échec(s)" % len(ECHECS))
    sys.exit(1)
print("Tout est vert.")
