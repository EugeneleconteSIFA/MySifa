"""
Export de l'inventaire matières (Excel et PDF).

Ce que le test verrouille :
- une laize sans stock ni emplacement ne fait pas de ligne ;
- le métrage d'une bobine = bobines × métrage standard de la fiche ;
- un emplacement « NC … » est compté hors stock disponible ;
- les filtres de l'écran (catégorie, recherche) s'appliquent ;
- le classeur a ses quatre feuilles, le PDF se construit.

Lancer : python3 tests/test_inventaire_matieres_export.py
"""

import io
import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path

RACINE = Path(__file__).resolve().parents[1]
if str(RACINE) not in sys.path:
    sys.path.insert(0, str(RACINE))
os.chdir(RACINE)
os.environ["DB_PATH"] = tempfile.NamedTemporaryFile(suffix=".db", delete=False).name
os.environ["ERP_MIRROR_DB"] = os.environ["DB_PATH"] + ".absent"
import database  # noqa: F401,E402  — le shim d'abord
from database import get_db  # noqa: E402
from app.routers.stock import matieres_inventaire_donnees  # noqa: E402
from app.services import inventaire_matieres_export as ime  # noqa: E402

FAIL = []


def check(label, got, expected):
    ok = got == expected
    print(("ok   " if ok else "KO   ") + label.ljust(60) + f"{got!r}"
          + ("" if ok else f"   attendu {expected!r}"))
    if not ok:
        FAIL.append(label)


with get_db() as c:
    mid = c.execute(
        """INSERT INTO matieres_premieres (categorie, reference, designation, actif, is_europe, prix_par_laize,
                                           suivi_bobine, metres_lineaires_par_bobine)
           VALUES ('frontal', 'TEST-EXP', 'Test Thermique export', 1, 0, 0, 0, 10000)""").lastrowid
    c.execute("INSERT INTO mp_stock (matiere_id, quantite) VALUES (?, 2.5)", (mid,))
    lz = []
    for v in (991.0, 992.0):
        lid = c.execute("INSERT INTO mp_laizes (valeur_mm, label, ordre, actif, created_at) VALUES (?, ?, 0, 1, 'x')",
                        (v, "%g mm" % v)).lastrowid
        c.execute("INSERT INTO mp_matiere_laizes (matiere_id, laize_id) VALUES (?, ?)", (mid, lid))
        lz.append(lid)
    c.execute("INSERT INTO mp_stock_laize (matiere_id, laize_id, quantite) VALUES (?, ?, 2.5)", (mid, lz[0]))
    c.execute("INSERT INTO mp_emplacements (matiere_id, laize_id, emplacement, quantite) VALUES (?, ?, 'A14', 2.5)",
              (mid, lz[0]))
    c.execute("INSERT INTO mp_emplacements (matiere_id, laize_id, emplacement, quantite) VALUES (?, ?, 'NC KL', 1)",
              (mid, lz[0]))

    tout = matieres_inventaire_donnees(c)
    m = next(x for x in tout if x["id"] == mid)
    check("laize sans stock ni emplacement écartée", [l["laize_mm"] for l in m["laizes"]], [991.0])
    check("métrage = bobines × métrage standard", (m["metres"], m["laizes"][0]["metres"]), (25000, 25000))
    check("non-conformité hors stock", m["nc"], 1.0)
    check("filtre catégorie", {x["categorie"] for x in matieres_inventaire_donnees(c, "frontal")}, {"frontal"})
    check("filtre recherche", [x["id"] for x in matieres_inventaire_donnees(c, q="thermique export")], [mid])

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(ime.classeur(tout, "", datetime.now())))
    check("classeur : quatre feuilles", wb.sheetnames, ["Synthèse", "Par référence", "Par laize", "Par emplacement"])
    empl = [r[0] for r in wb["Par emplacement"].iter_rows(min_row=4, values_only=True)
            if r[2] == "Test Thermique export"]
    check("une ligne par emplacement, triées", empl, ["A14", "NC KL"])
    check("PDF construit", ime.pdf(tout, "", datetime.now(), "test")[:4], b"%PDF")
    c.rollback()

print()
if FAIL:
    print("ÉCHEC : %d contrôle(s) en erreur" % len(FAIL))
    sys.exit(1)
print("Export inventaire matières : tout est vert.")
