"""Emplacements des matières premières et recherche d'articles RVGI (MyStock).

Ce que ces tests verrouillent :
- un emplacement se pose, se remplace, se retire, et ne touche jamais au stock
  total de la matière ;
- une matière laizée exige une laize qui lui est associée ;
- le code d'emplacement suit le format de la grille (lettre + chiffres) ;
- la liste des matières porte ses emplacements ;
- la recherche RVGI ignore accents et casse, reste dans le périmètre des
  réceptions et signale l'article déjà apparié.
"""

from __future__ import annotations

import os
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

_tmp = tempfile.mkdtemp(prefix="mysifa_mp_empl_test_")
_db = os.path.join(_tmp, "test.db")
_miroir = os.path.join(_tmp, "erp_mirror.db")
os.environ["DB_PATH"] = _db
os.environ["ERP_MIRROR_DB"] = _miroir

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

USER = {"id": 1, "nom": "Testeur", "email": "test@local", "role": "direction"}


def _construire_miroir():
    """Les colonnes que la recherche lit, pas le schéma complet de mat_mat."""
    c = sqlite3.connect(_miroir)
    c.executescript(
        """
        CREATE TABLE mat_mat (id INTEGER PRIMARY KEY, corbeille INTEGER, type INTEGER,
            code1 TEXT, code2 TEXT, libc1 TEXT, libt2 TEXT, ref TEXT);
        INSERT INTO mat_mat (corbeille,type,code1,code2,libc1,libt2,ref) VALUES
            (0,3,'982','0001','Vélin Torraspale - FSC','Bobine 12 000 ml',NULL),
            (0,6,'982','0001','Vélin hors périmètre','',NULL),
            (1,3,'982','0002','Vélin à la corbeille','',NULL),
            (0,13,'1101','0003','Tube Ø 76 mm, Lg 1.500 mm','Palette 338 tubes',NULL),
            (0,9,'897','0287','Cliché ronds 40 mm','',NULL);
        """
    )
    c.commit()
    c.close()


_construire_miroir()


def _client():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.routers.stock import router as stock_router

    app = FastAPI()
    app.include_router(stock_router)
    return TestClient(app, raise_server_exceptions=True)


class TestMpEmplacements(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import database  # noqa: F401 — le shim d'abord, voir CLAUDE.md
        from database import get_db

        with get_db() as conn:
            cur = conn.execute(
                "INSERT INTO matieres_premieres (categorie, reference, designation, actif) "
                "VALUES ('mandrin', 'T-EMPL-76', 'Tube test', 1)")
            cls.mandrin_id = cur.lastrowid
            conn.execute("INSERT INTO mp_stock (matiere_id, quantite) VALUES (?, 7)",
                         (cls.mandrin_id,))
            cur = conn.execute(
                "INSERT INTO matieres_premieres (categorie, sous_section, reference, designation, actif) "
                "VALUES ('frontal', 'Velin', 'T-EMPL-VEL', 'Vélin test', 1)")
            cls.frontal_id = cur.lastrowid
            conn.execute("INSERT INTO mp_stock (matiere_id, quantite) VALUES (?, 0)",
                         (cls.frontal_id,))
            laize = conn.execute("SELECT id FROM mp_laizes ORDER BY id LIMIT 2").fetchall()
            if len(laize) < 2:
                for v in (330, 250):
                    conn.execute(
                        "INSERT INTO mp_laizes (valeur_mm, label, ordre, actif) VALUES (?,?,?,1)",
                        (v, f"{v} mm", v))
                laize = conn.execute("SELECT id FROM mp_laizes ORDER BY id LIMIT 2").fetchall()
            cls.laize_ok, cls.laize_autre = int(laize[0]["id"]), int(laize[1]["id"])
            conn.execute("INSERT INTO mp_matiere_laizes (matiere_id, laize_id) VALUES (?,?)",
                         (cls.frontal_id, cls.laize_ok))
            conn.execute(
                "INSERT INTO erp_article_matiere (code1, code2, type_code, matiere_id, origine, "
                "created_at, created_by_name) VALUES ('1101','0003',15,?,'manuel','2026-09-25','t')",
                (cls.mandrin_id,))
            conn.commit()
        cls.client = _client()

    def setUp(self):
        self.patchs = [
            patch("app.routers.stock.get_current_user", return_value=dict(USER)),
            patch("app.routers.stock.user_has_app_access", return_value=True),
        ]
        for p in self.patchs:
            p.start()
        self.addCleanup(lambda: [p.stop() for p in self.patchs])

    def _poser(self, mid, **body):
        return self.client.post(f"/api/stock/matieres/{mid}/emplacements", json=body)

    def _stock(self, mid):
        from database import get_db
        with get_db() as conn:
            return conn.execute("SELECT quantite FROM mp_stock WHERE matiere_id=?",
                                (mid,)).fetchone()["quantite"]

    def test_poser_remplacer_retirer_sans_toucher_au_stock(self):
        r = self._poser(self.mandrin_id, emplacement="a121", quantite=3)
        self.assertEqual(r.status_code, 200, r.text)
        empl = r.json()["emplacements"]
        self.assertEqual([(e["emplacement"], e["quantite"]) for e in empl], [("A121", 3.0)])

        r = self._poser(self.mandrin_id, emplacement="A121", quantite=5)
        self.assertEqual([e["quantite"] for e in r.json()["emplacements"]], [5.0])
        self.assertEqual(self._stock(self.mandrin_id), 7)

        liste = self.client.get("/api/stock/matieres").json()
        m = next(x for x in liste if x["id"] == self.mandrin_id)
        self.assertEqual([e["emplacement"] for e in m["emplacements"]], ["A121"])

        eid = r.json()["emplacements"][0]["id"]
        r = self.client.delete(f"/api/stock/matieres/{self.mandrin_id}/emplacements/{eid}")
        self.assertEqual(r.json()["emplacements"], [])

    def test_quantite_zero_retire(self):
        self._poser(self.mandrin_id, emplacement="Z0", quantite=2)
        r = self._poser(self.mandrin_id, emplacement="Z0", quantite=0)
        self.assertEqual(r.json()["emplacements"], [])

    def test_format_emplacement(self):
        for code in ("", "121", "A-12", "AB12"):
            r = self._poser(self.mandrin_id, emplacement=code, quantite=1)
            self.assertEqual(r.status_code, 400, code)
        r = self._poser(self.mandrin_id, emplacement="A1", quantite=-1)
        self.assertEqual(r.status_code, 400)

    def test_laize_obligatoire_et_associee(self):
        self.assertEqual(self._poser(self.frontal_id, emplacement="B2", quantite=1).status_code, 400)
        r = self._poser(self.frontal_id, emplacement="B2", quantite=1, laize_id=self.laize_autre)
        self.assertEqual(r.status_code, 400)
        r = self._poser(self.frontal_id, emplacement="B2", quantite=4, laize_id=self.laize_ok)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["emplacements"][0]["laize_id"], self.laize_ok)

    def test_recherche_rvgi(self):
        r = self.client.get("/api/stock/matieres/rvgi-recherche", params={"q": "velin"})
        arts = r.json()["articles"]
        # mat_mat type 3 = ligne d'achat 5 (frontal Velin). L'article à la
        # corbeille ne sort pas, même si son libellé correspond.
        codes = {(a["code1"], a["code2"], a["type_code"]) for a in arts}
        self.assertIn(("982", "0001", 5), codes)
        self.assertNotIn(("982", "0002", 5), codes)
        a = next(a for a in arts if a["type_code"] == 5)
        self.assertEqual((a["categorie"], a["sous_section"]), ("frontal", "Velin"))

        arts = self.client.get("/api/stock/matieres/rvgi-recherche",
                               params={"q": "1101/0003"}).json()["articles"]
        self.assertEqual(len(arts), 1)
        self.assertEqual(arts[0]["matiere_id"], self.mandrin_id)

        arts = self.client.get("/api/stock/matieres/rvgi-recherche",
                               params={"q": "cliché"}).json()["articles"]
        self.assertEqual(arts, [])
        self.assertEqual(self.client.get("/api/stock/matieres/rvgi-recherche",
                                         params={"q": "v"}).json()["articles"], [])


def _lire_xlsx(contenu):
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(contenu)).active
    lignes = [list(r) for r in ws.iter_rows(values_only=True)]
    return lignes[0], lignes[1:]


class TestExportInventaire(unittest.TestCase):
    """Exports Excel des inventaires PF (par emplacement / par référence) et MP."""

    @classmethod
    def setUpClass(cls):
        import database  # noqa: F401
        from database import get_db

        now = "2026-09-25T10:00:00"
        with get_db() as conn:
            cur = conn.execute(
                "INSERT INTO produits (reference, designation, unite, created_at, updated_at) "
                "VALUES ('T-EXP-PF', 'Étiquette test', 'étiquette', ?, ?)", (now, now))
            pid = cur.lastrowid
            for empl, q, d in (("C11", 100, "2026-09-01T08:00:00"),
                               ("C11", 50, "2026-09-10T08:00:00"),
                               ("D22", 30, "2026-08-15T08:00:00")):
                conn.execute(
                    "INSERT INTO lots_stock (produit_id, emplacement, quantite_initiale, "
                    "quantite_restante, date_entree, created_at) VALUES (?,?,?,?,?,?)",
                    (pid, empl, q, q, d, now))
            # Lot épuisé : ne doit pas sortir.
            conn.execute(
                "INSERT INTO lots_stock (produit_id, emplacement, quantite_initiale, "
                "quantite_restante, date_entree, created_at) VALUES (?,'E33',10,0,?,?)",
                (pid, now, now))
            conn.execute(
                "INSERT INTO inventaires_sessions (emplacement, operateur_nom, date_validation) "
                "VALUES ('C11', 'Opérateur A', '2026-09-20T09:00:00')")
            conn.commit()
        cls.client = _client()

    def setUp(self):
        self.patchs = [
            patch("app.routers.stock.get_current_user", return_value=dict(USER)),
            patch("app.routers.stock.user_has_app_access", return_value=True),
        ]
        for p in self.patchs:
            p.start()
        self.addCleanup(lambda: [p.stop() for p in self.patchs])

    def test_pf_par_emplacement(self):
        r = self.client.get("/api/stock/inventaire-v2/export", params={"par": "emplacement"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertIn("attachment", r.headers["content-disposition"])
        entetes, lignes = _lire_xlsx(r.content)
        self.assertEqual(entetes[0], "Emplacement")
        self.assertEqual(entetes[-2:], ["Quantité comptée", "Écart"])
        mien = [l for l in lignes if l[2] == "T-EXP-PF"]
        self.assertEqual([(l[0], l[4]) for l in mien], [("C11", 150), ("D22", 30)])
        c11 = mien[0]
        self.assertEqual(c11[9], "Opérateur A")
        self.assertEqual(mien[1][8], "Jamais")
        self.assertTrue(str(c11[-1]).startswith("=IF("))

    def test_pf_par_reference(self):
        r = self.client.get("/api/stock/inventaire-v2/export", params={"par": "reference"})
        entetes, lignes = _lire_xlsx(r.content)
        self.assertEqual(entetes[0], "Référence")
        mien = [l for l in lignes if l[0] == "T-EXP-PF"]
        self.assertEqual(len(mien), 1)
        self.assertEqual(mien[0][2], 180)
        self.assertEqual(mien[0][4], 2)
        self.assertEqual(mien[0][5], "C11 : 150 · D22 : 30")
        # D22 n'a jamais été inventorié : la référence non plus, en partie.
        self.assertEqual(mien[0][7], "Jamais")

    def test_pf_par_invalide(self):
        r = self.client.get("/api/stock/inventaire-v2/export", params={"par": "lot"})
        self.assertEqual(r.status_code, 400)

    def test_mp_une_ligne_par_reference(self):
        from database import get_db
        with get_db() as conn:
            cur = conn.execute(
                "INSERT INTO matieres_premieres (categorie, reference, designation, actif) "
                "VALUES ('carton', 'T-EXP-CART', 'Carton export', 1)")
            mid = cur.lastrowid
            conn.execute("INSERT INTO mp_stock (matiere_id, quantite) VALUES (?, 12)", (mid,))
            conn.execute(
                "INSERT INTO mp_emplacements (matiere_id, laize_id, emplacement, quantite) "
                "VALUES (?, 0, 'F4', 5)", (mid,))
            conn.commit()
        r = self.client.get("/api/stock/matieres/inventaire/export")
        self.assertEqual(r.status_code, 200, r.text)
        entetes, lignes = _lire_xlsx(r.content)
        self.assertEqual(entetes[2], "Référence")
        mien = [l for l in lignes if l[2] == "T-EXP-CART"]
        self.assertEqual(len(mien), 1)
        self.assertEqual(mien[0][4], 12)
        self.assertEqual(mien[0][7], "F4 : 5")
        self.assertEqual(mien[0][8], "Jamais")
        self.assertEqual(mien[0][10], "À faire")
        refs = [l[2] for l in lignes]
        self.assertEqual(len(refs), len(set(refs)), "une ligne par référence")


if __name__ == "__main__":
    unittest.main()
